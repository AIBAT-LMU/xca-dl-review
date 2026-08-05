#!/usr/bin/env python3
"""
04_tab_s6_performance.py -- Table S6: Reported performance by metric
Output: results/tables/TableS6_performance.docx

Summarises each performance metric across the studies' PROPOSED models, scoped
to the tasks the metric belongs to (Dice/IoU -> segmentation; sensitivity/
specificity/AUC -> detection & classification; accuracy/F1 -> all). Per study
the best reported value is taken.

performance_* values are noisy free text -- table references ("0.73 (Table 1)"),
confidence intervals ("0.693 (0.685-0.701)"), per-bin/per-segment breakdowns
("1-49:0.29, ...") and cross-referenced multi-values. parse_pct() strips the
parentheticals and rejects the labelled/binned strings, so a naive min/max no
longer picks up a table number or a CI bound as a score. A handful of confirmed
mis-extractions are excluded explicitly.

Median [IQR] is the headline; min-max ranges are outlier-sensitive and are given
for completeness with the source study.
"""

import os
import re
import statistics

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

import common as C

OUT_PATH = os.path.join(C.OUT_TABLES, "TableS6_performance.docx")
NR = {"", "not_reported", "not reported", "n/a", "none", "unclear"}

# Confirmed mis-extractions, verified against the paper's own text.
#   ref 162 performance_iou 0.032 -- the proposed few-shot VOS method reports a
#   J&F score of 83.5%; the 0.032 was picked up from a "3.2.x" section number.
#   (ref is the current 1-196 numbering; was 182 before the renumber.)
EXCLUDE = {("performance_iou", 162)}


def parse_pct(s):
    s = (s or "").strip()
    if s.lower() in NR or ":" in s:          # binned / labelled / cross-ref
        return []
    s = re.sub(r"\([^)]*\)", " ", s)          # drop (Table N) and (CI a-b)
    out = []
    for tok in re.split(r"[|/,]", s):
        tok = tok.split("±")[0].strip()
        for nstr in re.findall(r"\d+\.?\d*", tok)[:2]:   # 'a-b' range -> both
            v = float(nstr)
            if v <= 1.5:
                v *= 100
            if 0 <= v <= 100:
                out.append(v)
    return out


conn = C.connect()
papers = C.load_papers(conn)
ref_of = {p.paper_id: p.ref for p in papers}
by_domain, _ = C.application_domains(conn, papers)

proposed = {}
for pid, fn, v in conn.execute(
        "SELECT paper_id, field_name, value FROM extractions "
        "WHERE field_name LIKE 'model_role__%' AND value='Proposed'"):
    proposed.setdefault(pid, []).append(int(re.search(r"__(\d+)$", fn).group(1)))


def study_best(pid, base):
    if (base, ref_of[pid]) in EXCLUDE:
        return None
    vals = []
    for i in proposed.get(pid, []):
        r = conn.execute("SELECT value FROM extractions WHERE paper_id=? "
                         "AND field_name=?", (pid, f"{base}__{i}")).fetchone()
        if r and r[0]:
            vals += parse_pct(r[0])
    return max(vals) if vals else None


SEG = by_domain["Vessel segmentation"]
DETCLS = (by_domain["Stenosis detection"] | by_domain["Other detection"]
          | by_domain["Classification"])
ALL = {p.paper_id for p in papers}
METRICS = [
    ("Dice (segmentation)", "performance_dice", SEG),
    ("IoU (segmentation)", "performance_iou", SEG),
    ("Sensitivity (detection/classification)", "performance_recall", DETCLS),
    ("Specificity (detection/classification)", "performance_specificity", DETCLS),
    ("AUC-ROC (detection/classification)", "performance_aucroc", DETCLS),
    ("Accuracy (all tasks)", "performance_accuracy", ALL),
    ("F1 (all tasks)", "performance_f1", ALL),
]

rows = []
print("Table S6: Reported performance by metric")
for label, base, scope in METRICS:
    best = {pid: study_best(pid, base) for pid in scope}
    best = {pid: v for pid, v in best.items() if v is not None}
    vals = sorted(best.values())
    q1, med, q3 = (statistics.quantiles(vals, n=4)[0], statistics.median(vals),
                   statistics.quantiles(vals, n=4)[2])
    lo = min(best.items(), key=lambda kv: kv[1])
    hi = max(best.items(), key=lambda kv: kv[1])
    rng = f"{lo[1]:.1f}–{hi[1]:.1f}"
    rowrefs = f"low {lo[1]:.1f}% [{ref_of[lo[0]]}], high {hi[1]:.1f}% [{ref_of[hi[0]]}]"
    rows.append((label, str(len(vals)), f"{med:.0f} [{q1:.0f}–{q3:.0f}]", rng, rowrefs))
    print(f"  {label:42s} n={len(vals):3d}  median {med:.0f}% [IQR {q1:.0f}–{q3:.0f}]  "
          f"range {rng}%  ({rowrefs})")


# ── Meta-analytical comparison on shared public benchmarks ──
# Corpus-wide pooling is precluded by heterogeneity, but studies that evaluate
# on the same public benchmark are comparable. eval_dataset__N and
# performance_*__N share the model_performance__N group, so each score links to
# its evaluation dataset. Counts differ from Table S4 (which counts studies that
# *use* a dataset): a study may use a benchmark only for training, or report a
# different metric on it -- e.g. CADICA is a detection set, so few report Dice.
# Dice is a segmentation metric, so this panel covers the shared *segmentation*
# benchmarks only. Detection sets (CADICA) are compared on detection metrics
# elsewhere, not here.
BENCHMARKS = {"DCA1": r"\bdca[-\s]?1\b", "ARCADE": r"arcade",
              "XCAD": r"xcad", "CHUAC": r"chuac"}
ed, dice_on, iou_on = {}, {}, {}
for pid, fn, v in conn.execute(
        "SELECT paper_id, field_name, value FROM extractions WHERE "
        "field_name LIKE 'eval_dataset__%' OR field_name LIKE 'performance_dice__%' "
        "OR field_name LIKE 'performance_iou__%'"):
    idx = re.search(r"__(\d+)$", fn).group(1)
    if fn.startswith("eval_dataset"):
        ed[(pid, idx)] = (v or "").strip()
    elif fn.startswith("performance_dice") and parse_pct(v):
        dice_on[(pid, idx)] = parse_pct(v)[0]
    elif fn.startswith("performance_iou") and parse_pct(v):
        iou_on[(pid, idx)] = True
uses = {name: {pid for pid, fn, v in conn.execute(
            "SELECT paper_id, field_name, value FROM extractions "
            "WHERE field_name LIKE 'dataset_name__%'")
        if re.search(pat, v or "", re.I)} for name, pat in BENCHMARKS.items()}
conn.close()

print("\nMeta-analytical comparison on shared benchmarks "
      "(uses = Table S4; Dice-on = evaluated + Dice reported):")
print(f"  {'benchmark':10s}{'uses':>6}{'eval-on':>9}{'Dice-on':>9}{'IoU-on':>8}  Dice median [range]")
bench_rows = []
for name, pat in BENCHMARKS.items():
    evalon = {k[0] for k, d in ed.items() if re.search(pat, d, re.I)}
    dstud = {}
    for k, val in dice_on.items():
        if k in ed and re.search(pat, ed[k], re.I):
            dstud[k[0]] = max(dstud.get(k[0], 0), val)
    ion = {k[0] for k in iou_on if k in ed and re.search(pat, ed[k], re.I)}
    ds = sorted(dstud.values())
    stat = (f"{statistics.median(ds):.0f} [{min(ds):.1f}–{max(ds):.1f}]"
            if len(ds) >= 3 else "—")
    bench_rows.append((name, str(len(uses[name])), str(len(evalon)),
                       str(len(dstud)), str(len(ion)), stat))
    print(f"  {name:10s}{len(uses[name]):>6}{len(evalon):>9}{len(dstud):>9}{len(ion):>8}  {stat}")


C.ensure_dirs()
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(2)
section.left_margin = section.right_margin = Cm(1.5)

p = doc.add_paragraph()
run = p.add_run("Table S6. ")
run.bold = True
run.font.size = Pt(11)
run.font.name = "Arial"
run = p.add_run("Reported performance by metric")
run.font.size = Pt(11)
run.font.name = "Arial"

table = doc.add_table(rows=len(rows) + 1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
table.style = "Table Grid"
for i, text in enumerate(["Metric", "n", "Median % [IQR]", "Range % (source study)"]):
    cell = table.rows[0].cells[i]
    cell.text = ""
    r = cell.paragraphs[0].add_run(text)
    r.bold = True
    r.font.size = Pt(9)
    r.font.name = "Arial"
    tcPr = cell._element.get_or_add_tcPr()
    tcPr.append(tcPr.makeelement(qn("w:shd"),
                                 {qn("w:fill"): "E8E8E8", qn("w:val"): "clear"}))
for i, (label, n, medi, rng, rowrefs) in enumerate(rows):
    for j, text in enumerate((label, n, medi, f"{rng} ({rowrefs})")):
        cell = table.rows[i + 1].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(text)
        r.font.size = Pt(8) if j == 3 else Pt(9)
        r.font.name = "Arial"
for row in table.rows:
    row.cells[0].width = Cm(5)
    row.cells[1].width = Cm(1)
    row.cells[2].width = Cm(2.5)
    row.cells[3].width = Cm(9.5)

p = doc.add_paragraph()
run = p.add_run(
    "Values are the best reported figure for each study's proposed model(s). "
    "Ranges are outlier-sensitive; median [IQR] is the more robust summary.")
run.italic = True
run.font.size = Pt(8)
run.font.name = "Arial"

# ── Second panel: benchmark comparison ──
p = doc.add_paragraph()
run = p.add_run("Reported Dice on shared segmentation benchmarks")
run.bold = True
run.font.size = Pt(10)
run.font.name = "Arial"

bt = doc.add_table(rows=len(bench_rows) + 1, cols=6)
bt.alignment = WD_TABLE_ALIGNMENT.LEFT
bt.style = "Table Grid"
for i, text in enumerate(["Benchmark", "Uses ᵃ", "Eval-on", "Dice-on",
                          "IoU-on", "Dice median [range] %"]):
    cell = bt.rows[0].cells[i]
    cell.text = ""
    r = cell.paragraphs[0].add_run(text)
    r.bold = True
    r.font.size = Pt(9)
    r.font.name = "Arial"
    tcPr = cell._element.get_or_add_tcPr()
    tcPr.append(tcPr.makeelement(qn("w:shd"),
                                 {qn("w:fill"): "E8E8E8", qn("w:val"): "clear"}))
for i, brow in enumerate(bench_rows):
    for j, text in enumerate(brow):
        cell = bt.rows[i + 1].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(text)
        r.font.size = Pt(9)
        r.font.name = "Arial"

p = doc.add_paragraph()
run = p.add_run(
    "ᵃ Studies using the dataset (Table S4); the count falls to Dice-reporting "
    "because a study may use a benchmark only for training or report IoU rather "
    "than Dice. Even on DCA1, a single binary-segmentation task, reported Dice "
    "spans 61.5–93.9%, so pooling is precluded by protocol variation, not only "
    "differing datasets.")
run.italic = True
run.font.size = Pt(8)
run.font.name = "Arial"

doc.save(OUT_PATH)
print(f"\nSaved → {OUT_PATH}")
