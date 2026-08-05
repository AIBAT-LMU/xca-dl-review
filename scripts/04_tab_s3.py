#!/usr/bin/env python3
"""
04_tab_s3.py -- Table S3: Training data, annotation and cohort characteristics
Output: results/tables/TableS3_training.docx

Datasets are a repeat group (dataset_*__1..N) with a role; the training set is
the Training / Training+Test instance (a single split reported as one size, so
the size reflects the whole dataset, not the training portion alone). Centre is
the paper-level multi_center field.
"""

import os
import re
import statistics
from collections import defaultdict

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

import common as C

OUT_PATH = os.path.join(C.OUT_TABLES, "TableS3_training.docx")

SAMPLE_BINS = [("≤100", 0, 100), ("101–500", 101, 500),
               ("501–1,000", 501, 1000), ("1,001–5,000", 1001, 5000),
               ("5,001–10,000", 5001, 10000), (">10,000", 10001, 1e12)]
PATIENT_BINS = [("≤50", 0, 50), ("51–100", 51, 100),
                ("101–200", 101, 200), ("201–500", 201, 500),
                ("501–1,000", 501, 1000), (">1,000", 1001, 1e12)]
EXPERT_KW = ("cardiolog", "radiolog", "expert", "clinician", "physician",
             "specialist", "experienced", "senior", "attending")
STUDENT_KW = ("student", "trainee", "resident", "junior", "intern")
NR = {"", "not_reported", "not reported", "n/a", "none", "unclear"}


def parse_num(s):
    if not s:
        return None
    s = s.strip()
    if s.lower() in NR:
        return None
    m = re.search(r"[\d][\d,\.]*", s.replace(" ", ""))
    if not m:
        return None
    try:
        return float(m.group().replace(",", ""))
    except ValueError:
        return None


conn = C.connect()
papers = C.load_papers(conn)
N = len(papers)
ref_of = {p.paper_id: p.ref for p in papers}

# Repeat-group dataset instances: paper_id -> idx -> {field: value}
inst = defaultdict(lambda: defaultdict(dict))
for pid, fname, val in conn.execute(
        "SELECT paper_id, field_name, value FROM extractions "
        "WHERE field_name LIKE 'dataset_%__%' AND value IS NOT NULL AND value!=''"):
    m = re.match(r"(dataset_\w+)__(\d+)$", fname)
    if m:
        inst[pid][int(m.group(2))][m.group(1)] = val.strip()


def scalar(field):
    return {pid: (v or "").strip() for pid, v in conn.execute(
        "SELECT paper_id, value FROM extractions WHERE field_name=?", (field,))}


def repeated(base):
    out = defaultdict(list)
    for pid, fname, val in conn.execute(
            "SELECT paper_id, field_name, value FROM extractions "
            "WHERE field_name LIKE ? ESCAPE '\\' AND value IS NOT NULL AND value!=''",
            (base + "\\_\\_%",)):
        out[pid].append(val.strip())
    return out


# Training-set sizes: largest Training / Training+Test instance per paper.
train_samples, train_patients = {}, {}
for pid, idxs in inst.items():
    ss, pp = [], []
    for f in idxs.values():
        if "Training" in f.get("dataset_role", ""):
            v = parse_num(f.get("dataset_size_samples", ""))
            if v:
                ss.append(v)
            w = parse_num(f.get("dataset_size_patients", ""))
            if w:
                pp.append(w)
    if ss:
        train_samples[pid] = max(ss)
    if pp:
        train_patients[pid] = max(pp)

center = scalar("multi_center")
expertise = repeated("annotator_expertise")
ann_count = repeated("annotator_count")
iaa = repeated("inter_annotator_agreement")
age, sex = scalar("age_reported"), scalar("sex_reported")
presentation = scalar("patient_presentation")
conn.close()


rows = []  # (label, n, pct, refs, indent, header)


def add(members, label, indent=1):
    refs = sorted(ref_of[p] for p in members)
    rows.append((label, len(members), C.pct(len(members), N),
                 C.ref_ranges(refs), indent, False))


def head(label):
    rows.append((label, "", "", "", 0, True))


def bin_rows(values, bins):
    for label, lo, hi in bins:
        members = [p for p, v in values.items() if v is not None and lo <= v <= hi]
        add(members, label)
    nr = [p.paper_id for p in papers if values.get(p.paper_id) is None]
    add(nr, "Not reported")


def stat_line(values, unit):
    vs = sorted(v for v in values.values() if v)
    if not vs:
        return
    q1, q3 = statistics.quantiles(vs, n=4)[0], statistics.quantiles(vs, n=4)[2]
    rows.append((f"median {statistics.median(vs):,.0f} {unit} "
                 f"(IQR {q1:,.0f}–{q3:,.0f}; range {min(vs):,.0f}–{max(vs):,.0f}; "
                 f"n={len(vs)})", "", "", "", 1, False))


# ── Data source (centre) ──
head(f"Data source (N={N})")
add([p for p, v in center.items() if v == "Single-center"], "Single-centre")
add([p for p, v in center.items() if v == "Multi-center"], "Multi-centre")
add([p for p, v in center.items() if v == "Unclear"], "Unclear")

# ── Data provenance (country) ──
COUNTRY_MERGE = {"republic of korea": "South Korea", "korea": "South Korea",
                 "united kingdom": "UK", "uk": "UK", "usa": "USA", "us": "USA"}
NR_COUNTRY = {"not reported", "not_reported", "n/a", "unclear", "",
              "various", "multiple", "multi-national"}
country = defaultdict(set)   # country -> {paper_id}; a study may use several
for pid, idxs in inst.items():
    for f in idxs.values():
        cl = (f.get("dataset_country", "") or "").strip().lower()
        if cl and cl not in NR_COUNTRY:
            country[COUNTRY_MERGE.get(cl, cl.title())].add(pid)
reported_country = set().union(*country.values()) if country else set()
head(f"Data provenance (country, N={N}) ᵈ")
TOP = ["China", "Mexico", "Russia", "South Korea", "USA"]
for c in TOP:
    add(country.get(c, set()), c)
add(reported_country - set().union(*(country.get(c, set()) for c in TOP)), "Other country")
add([p.paper_id for p in papers if p.paper_id not in reported_country], "Not reported")

# ── Training set size ──
head(f"Training set size, samples (N={N}) ᵃ")
stat_line(train_samples, "samples")
bin_rows(train_samples, SAMPLE_BINS)

head(f"Patient count (N={N}) ᵇ")
stat_line(train_patients, "patients")
bin_rows(train_patients, PATIENT_BINS)

# ── Annotation ──
head(f"Annotation (N={N})")
described = {p for p, vs in expertise.items() if any(v.lower() not in NR for v in vs)}
expert = {p for p, vs in expertise.items()
          if any(k in v.lower() for v in vs for k in EXPERT_KW)}
student = {p for p, vs in expertise.items()
           if any(k in v.lower() for v in vs for k in STUDENT_KW)}
two_plus = {p for p, vs in ann_count.items()
            if any((parse_num(v) or 0) >= 2 for v in vs)}
iaa_rep = {p for p, vs in iaa.items() if any(v.lower() not in NR for v in vs)}
add(described, "Annotator expertise described")
add(expert, "Expert (cardiologist/radiologist) ᶜ", 2)
add(student, "Student/trainee", 2)
add(two_plus, "≥2 annotators")
add(iaa_rep, "Inter-annotator agreement reported")

# ── Cohort characterisation ──
head(f"Cohort characterisation (N={N})")
add([p for p, v in age.items() if v.upper() == "Y"], "Patient age reported")
add([p for p, v in sex.items() if v.upper() == "Y"], "Patient sex reported")
add([p for p, v in presentation.items() if v.strip().lower() not in NR],
    "Clinical presentation reported")


# ── Console ──
print(f"Table S3: Training data characteristics (N={N})")
for label, n, pct, refs, indent, hdr in rows:
    pfx = "  " * indent
    print(f"\n{pfx}{label}" if hdr else f"{pfx}{label}: {n} ({pct}%)" if n != ""
          else f"{pfx}{label}")


# ── DOCX ──
C.ensure_dirs()
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(2)
section.left_margin = section.right_margin = Cm(1.5)

p = doc.add_paragraph()
run = p.add_run("Table S3. ")
run.bold = True
run.font.size = Pt(11)
run.font.name = "Arial"
run = p.add_run("Training data, annotation and cohort characteristics")
run.font.size = Pt(11)
run.font.name = "Arial"

table = doc.add_table(rows=len(rows) + 1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
table.style = "Table Grid"
for i, text in enumerate(["Characteristic", "n", "%", "References"]):
    cell = table.rows[0].cells[i]
    cell.text = ""
    r = cell.paragraphs[0].add_run(text)
    r.bold = True
    r.font.size = Pt(9)
    r.font.name = "Arial"
    tcPr = cell._element.get_or_add_tcPr()
    tcPr.append(tcPr.makeelement(qn("w:shd"),
                                 {qn("w:fill"): "E8E8E8", qn("w:val"): "clear"}))

for i, (label, n, pct, refs, indent, hdr) in enumerate(rows):
    cell = table.rows[i + 1].cells[0]
    cell.text = ""
    r = cell.paragraphs[0].add_run("   " * indent + label)
    r.bold = hdr
    r.italic = indent == 1 and n == ""   # median/IQR lines
    r.font.size = Pt(9)
    r.font.name = "Arial"
    for j, text in enumerate((str(n) if n != "" else "",
                              str(pct) if pct != "" else "", refs)):
        cell = table.rows[i + 1].cells[j + 1]
        cell.text = ""
        r = cell.paragraphs[0].add_run(text)
        r.font.size = Pt(8) if j == 2 else Pt(9)
        r.font.name = "Arial"

for row in table.rows:
    row.cells[0].width = Cm(5.5)
    row.cells[1].width = Cm(1)
    row.cells[2].width = Cm(1)
    row.cells[3].width = Cm(10.5)

p = doc.add_paragraph()
run = p.add_run(
    "ᵃ Training or training-plus-test set; sizes are heterogeneous (images, "
    "frames, cine loops or patches) and indicative, not directly comparable. "
    "ᵇ Unique patients. ᶜ Cardiologist, radiologist or experienced clinician. "
    "ᵈ A study may use datasets from more than one country, so country rows can "
    "exceed the total. Reference numbers correspond to the supplementary "
    "reference list L1.")
run.italic = True
run.font.size = Pt(8)
run.font.name = "Arial"

doc.save(OUT_PATH)
print(f"\nSaved → {OUT_PATH}")
