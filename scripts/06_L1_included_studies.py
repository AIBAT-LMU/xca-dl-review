#!/usr/bin/env python3
"""
06_L1_included_studies.py -- Supplementary Reference List L1: the 196 included
studies, numbered 1-196 (the ground-truth numbering used by every table/figure).

Source of citation text is xca_dl_review_references.bib (clean UTF-8; the pasted
zotero_list.txt lost its non-ASCII characters). Ordering and numbering come from
the database ref_number (1-196, Zotero AMA-alphabetical), matched to the bib by
DOI, so L1 and the tables cannot drift. AMA-ish output:

    N. Author AB, Author CD. Title. Journal. Year;Vol(Issue):Pages. doi:...
    N. Author AB. Title. In: Booktitle. Year. doi:...

Outputs results/tables/L1_included_studies.docx and .txt.
"""
import os
import re

from docx import Document
from docx.shared import Pt, Cm

import common as C

OUT_DOCX = os.path.join(C.OUT_TABLES, "L1_included_studies.docx")
OUT_TXT = os.path.join(C.OUT_TABLES, "L1_included_studies.txt")


def _field(body, key):
    """Brace-balanced value of `key = {...}` (handles nested braces in titles)."""
    m = re.search(rf"\b{key}\s*=\s*\{{", body, re.I)
    if not m:
        # quoted or bare value
        m2 = re.search(rf'\b{key}\s*=\s*"(.*?)"', body, re.I | re.S)
        if m2:
            return m2.group(1)
        m3 = re.search(rf"\b{key}\s*=\s*([^,\n}}]+)", body, re.I)
        return m3.group(1).strip() if m3 else None
    i = m.end()
    depth, start = 1, i
    while i < len(body) and depth:
        if body[i] == "{":
            depth += 1
        elif body[i] == "}":
            depth -= 1
        i += 1
    return body[start:i - 1]


def _clean(s):
    if not s:
        return ""
    s = s.replace("{", "").replace("}", "").replace("\\&", "&")
    return re.sub(r"\s+", " ", s).strip()


def _pages(s):
    s = _clean(s).replace("--", "-")
    m = re.match(r"^(\S+)-(\S+)$", s)
    if m and m.group(1) == m.group(2):   # "1186-1186" -> "1186"
        return m.group(1)
    return s


def _authors(raw):
    """'Last, First M and ...' -> 'Last FM, ...' (AMA; >6 authors -> first 3 et al)."""
    people = [p.strip() for p in re.split(r"\s+and\s+", _clean(raw)) if p.strip()]
    out = []
    for p in people:
        if "," in p:
            last, given = p.split(",", 1)
        else:  # 'First Last' fallback
            parts = p.split()
            last, given = parts[-1], " ".join(parts[:-1])
        initials = "".join(tok[0] for tok in given.split() if tok)
        out.append(f"{last.strip()} {initials}".strip())
    if len(out) > 6:
        return ", ".join(out[:3]) + ", et al"
    return ", ".join(out)


def parse_bib(path):
    text = open(path, encoding="utf-8").read()
    by_doi = {}
    for m in re.finditer(r"@(\w+)\{([^,]+),(.*?)\n\}", text, re.S):
        etype, body = m.group(1).lower(), m.group(3)
        doi = C.normalise_doi(_field(body, "doi") or "")
        if not doi:
            continue
        by_doi[doi] = dict(
            etype=etype,
            author=_authors(_field(body, "author") or ""),
            title=_clean(_field(body, "title")).rstrip(". "),
            journal=_clean(_field(body, "journal")),
            booktitle=_clean(_field(body, "booktitle")),
            year=_clean(_field(body, "year")),
            volume=_clean(_field(body, "volume")),
            number=_clean(_field(body, "number")),
            pages=_pages(_field(body, "pages")),
            doi=doi,
        )
    return by_doi


def citation(e):
    head = f"{e['author']}. {e['title']}."
    if e["etype"] == "inproceedings" or (e["booktitle"] and not e["journal"]):
        venue = f" In: {e['booktitle']}." if e["booktitle"] else ""
        yr = f" {e['year']}." if e["year"] else ""
        return f"{head}{venue}{yr} doi:{e['doi']}"
    j = f" {e['journal']}." if e["journal"] else ""
    yv = ""
    if e["year"]:
        yv = f" {e['year']}"
        if e["volume"]:
            yv += f";{e['volume']}"
            if e["number"]:
                yv += f"({e['number']})"
            if e["pages"]:
                yv += f":{e['pages']}"
        yv += "."
    return f"{head}{j}{yv} doi:{e['doi']}"


conn = C.connect()
papers = C.load_papers(conn)
conn_doi = {C.normalise_doi(r[0]): r[1] for r in conn.execute("SELECT doi, ref_number FROM papers")}
conn.close()

by_doi = parse_bib(C.BIB_PATH)

# Build L1 in ref_number order (1-196), verify every study matches a bib entry.
doi_of_ref = {v: k for k, v in conn_doi.items()}
lines, missing = [], []
for ref in range(1, len(papers) + 1):
    doi = doi_of_ref.get(ref)
    e = by_doi.get(doi)
    if e is None:
        missing.append((ref, doi))
        continue
    lines.append((ref, citation(e)))

if missing:
    for ref, doi in missing:
        print(f"  ! ref {ref} (doi {doi}) has no bib entry")
    raise SystemExit(f"{len(missing)} studies missing from the bib; L1 not written")

print(f"L1: {len(lines)} studies formatted (refs 1-{len(lines)}), all matched to bib by DOI.")

# ── TXT ──
C.ensure_dirs()
with open(OUT_TXT, "w", encoding="utf-8") as fh:
    for ref, cit in lines:
        fh.write(f"{ref}. {cit}\n")

# ── DOCX ──
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2)
sec.left_margin = sec.right_margin = Cm(2)
h = doc.add_paragraph()
r = h.add_run("Supplementary Reference List L1. ")
r.bold = True
r.font.size = Pt(11)
r.font.name = "Arial"
r = h.add_run("Included studies (n = 196).")
r.font.size = Pt(11)
r.font.name = "Arial"
for ref, cit in lines:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(f"{ref}. {cit}")
    run.font.size = Pt(9)
    run.font.name = "Arial"
doc.save(OUT_DOCX)
print(f"Saved -> {OUT_DOCX}\nSaved -> {OUT_TXT}")
