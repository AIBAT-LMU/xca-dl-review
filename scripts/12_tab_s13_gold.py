#!/usr/bin/env python3
"""
12_tab_s13_gold.py -- Table S13: human gold-standard comparison of the
LLM-assisted CLAIM 2024 and QUADAS-2 appraisal.
Output: results/tables/TableS13_gold_standard.docx

Two authors (ML background, C.G.; clinical background, V.O.I.) independently
rated a second random 10% subsample (n = 20) on the full QUADAS-2 (11
signalling questions, 3 applicability judgements) and 18 judgement-dependent
CLAIM 2024 items, from the same written guidance as the model and blind to its
output. Ratings are read from gold_standard_ratings.xlsx (sheet "ratings": ref,
rater, instrument, item, value; raters labelled "clinical" and "ml"), the
anonymised long-format export of the rating forms.

Reported:
  A. pairwise agreement by instrument: raw %, Gwet's AC1, Cohen's kappa
  B. per-item raw agreement for the three pairs
  C. aggregate figures per rater: CLAIM compliance on the 18 items,
     QUADAS-2 domains rated low risk, applicability rated low concern

Gwet's AC1 is reported alongside kappa because several items have a
near-degenerate marginal distribution (kappa paradox; see Methods).
"""

import math
import os
import re
import statistics as st
from collections import Counter, defaultdict

import openpyxl
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

import common as C

XLSX = os.environ.get("XCA_GOLD", os.path.join(os.path.dirname(C.DB_PATH),
                                                "gold_standard_ratings.xlsx"))
OUT_PATH = os.path.join(C.OUT_TABLES, "TableS13_gold_standard.docx")
RATERS = ("clinical", "ml", "LLM")
PAIRS = (("clinical", "ml"), ("clinical", "LLM"), ("ml", "LLM"))

### Read the ratings
wb = openpyxl.load_workbook(XLSX, data_only=True)
rows = [r for r in wb["ratings"].iter_rows(min_row=2, values_only=True) if r[0] is not None]
items = []
for _, _, inst, it, _ in rows:
    if (inst, str(it)) not in items:
        items.append((inst, str(it)))
refs = sorted({int(r[0]) for r in rows})


def norm(v):
    v = (v or "").strip().lower()
    return {"not_applicable": "na", "n/a": "na", "not applicable": "na",
            "low concern": "low", "high concern": "high"}.get(v, v)


def group(it):
    if it[0] == "CLAIM":
        return "CLAIM 2024"
    return "QUADAS-2 applicability" if it[1].endswith("app") else "QUADAS-2 signalling"


R = defaultdict(dict)   # R[rater][(ref, item)] = value
for ref, rater, inst, it, v in rows:
    R[rater][(int(ref), (inst, str(it)))] = norm(v)
conn = C.connect()
pid = {r[1]: r[0] for r in conn.execute("SELECT id, ref_number FROM papers")}
for ref in refs:
    for cn, it, res in conn.execute("SELECT checklist_name, item_id, result FROM checklist_results "
                                    "WHERE paper_id=?", (pid[ref],)):
        key = ("CLAIM" if cn.startswith("CLAIM") else "QUADAS-2", it)
        if key in items:
            R["LLM"][(ref, key)] = norm(res)
conn.close()
missing = [(r, ref, k) for r in RATERS for ref in refs for k in items if (ref, k) not in R[r]]
assert not missing, f"missing ratings: {missing[:5]}"


### Agreement statistics
def ac1(pairs):
    cats = sorted({a for a, _ in pairs} | {b for _, b in pairs})
    n = len(pairs)
    pa = sum(a == b for a, b in pairs) / n
    pi = {c: (sum(a == c for a, _ in pairs) + sum(b == c for _, b in pairs)) / (2 * n) for c in cats}
    q = len(cats)
    pe = sum(p * (1 - p) for p in pi.values()) / (q - 1) if q > 1 else 0.0
    return (pa - pe) / (1 - pe) if pe < 1 else float("nan")


def kappa(pairs):
    cats = sorted({a for a, _ in pairs} | {b for _, b in pairs})
    n = len(pairs)
    pa = sum(a == b for a, b in pairs) / n
    pe = sum((sum(a == c for a, _ in pairs) / n) * (sum(b == c for _, b in pairs) / n) for c in cats)
    return (pa - pe) / (1 - pe) if pe < 1 else float("nan")


def pairs_for(a, b, keys):
    return [(R[a][k], R[b][k]) for k in keys]


GROUPS = ("CLAIM 2024", "QUADAS-2 signalling", "QUADAS-2 applicability")
partA = []
for grp in GROUPS:
    keys = [(ref, it) for ref in refs for it in items if group(it) == grp]
    row = [grp, str(len(keys))]
    for a, b in PAIRS:
        p = pairs_for(a, b, keys)
        row.append(f"{100 * sum(x == y for x, y in p) / len(p):.0f}% / {ac1(p):.2f} / {kappa(p):.2f}")
    partA.append(row)

partB = []
for it in items:
    keys = [(ref, it) for ref in refs]
    row = [f"{it[0]} {it[1]}"]
    for a, b in PAIRS:
        p = pairs_for(a, b, keys)
        row.append(f"{100 * sum(x == y for x, y in p) / len(p):.0f}%")
    partB.append(row)

ROB = {"Patient selection": ["1.1", "1.2", "1.3"], "Index test": ["2.1", "2.2"],
       "Reference standard": ["3.1", "3.2"], "Flow and timing": ["4.1", "4.2", "4.3", "4.4"]}


def derive(vals):
    if any(v == "no" for v in vals):
        return "high"
    if all(v == "yes" for v in vals):
        return "low"
    return "unclear"


def compliance(rater, ref):
    y = sum(R[rater][(ref, it)] == "yes" for it in items if it[0] == "CLAIM")
    n = sum(R[rater][(ref, it)] == "no" for it in items if it[0] == "CLAIM")
    return 100 * y / (y + n) if y + n else float("nan")


partC = []
row = ["CLAIM 2024 compliance on the 18 items, mean (median)"]
for r in RATERS:
    cs = [compliance(r, ref) for ref in refs]
    row.append(f"{st.mean(cs):.0f}% ({st.median(cs):.0f}%)")
partC.append(row)
for dom, qs in ROB.items():
    row = [f"QUADAS-2 {dom.lower()}: studies rated low / unclear / high risk"]
    for r in RATERS:
        c = Counter(derive([R[r][(ref, ("QUADAS-2", q))] for q in qs]) for ref in refs)
        row.append(f"{c['low']} / {c['unclear']} / {c['high']}")
    partC.append(row)
row = ["QUADAS-2 domains rated low risk (of 80)"]
for r in RATERS:
    t = sum(derive([R[r][(ref, ("QUADAS-2", q))] for q in qs]) == "low" for ref in refs for qs in ROB.values())
    row.append(f"{t} ({100 * t / 80:.0f}%)")
partC.append(row)
row = ["QUADAS-2 applicability rated low concern (of 60)"]
for r in RATERS:
    t = sum(R[r][(ref, ("QUADAS-2", a))] == "low" for ref in refs for a in ("1.app", "2.app", "3.app"))
    row.append(f"{t} ({100 * t / 60:.0f}%)")
partC.append(row)

### Console summary
print(f"Table S13: gold standard, n = {len(refs)} studies, {len(items)} items")
for row in partA:
    print("  " + " | ".join(row))
for row in partC[-2:]:
    print("  " + " | ".join(row))

### DOCX
C.ensure_dirs()
doc = Document()


def shade(cell, fill):
    tcPr = cell._element.get_or_add_tcPr()
    tcPr.append(tcPr.makeelement(qn("w:shd"), {qn("w:fill"): fill, qn("w:val"): "clear"}))


def put(cell, text, *, bold=False, size=8, align="center", fill=None):
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = {"center": 1, "left": 0}[align]
    run = para.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Arial Narrow"
    if fill:
        shade(cell, fill)


def heading(text_bold, text):
    p = doc.add_paragraph()
    r = p.add_run(text_bold)
    r.bold = True
    r.font.size = Pt(10)
    r.font.name = "Arial"
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.name = "Arial"


def table(head, rows, widths, left_cols=(0,)):
    t = doc.add_table(rows=len(rows) + 1, cols=len(head))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    t.autofit = False
    for j, h in enumerate(head):
        put(t.rows[0].cells[j], h, bold=True, fill="F2F2F2", align="left" if j in left_cols else "center")
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            put(t.rows[i + 1].cells[j], v, align="left" if j in left_cols else "center")
    for row in t.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
    doc.add_paragraph()


p = doc.add_paragraph()
r = p.add_run("Table S13. ")
r.bold = True
r.font.size = Pt(11)
r.font.name = "Arial"
r = p.add_run(f"Human gold-standard comparison of the LLM-assisted appraisal (n = {len(refs)} studies)")
r.font.size = Pt(11)
r.font.name = "Arial"

heading("A. ", "Pairwise agreement by instrument: raw agreement / Gwet's AC1 / Cohen's kappa")
table(["Instrument", "Ratings", "Clinical vs ML rater", "Clinical rater vs model", "ML rater vs model"],
      partA, [4.2, 1.6, 3.4, 3.6, 3.4])

heading("B. ", "Per-item raw agreement")
table(["Item", "Clinical vs ML rater", "Clinical rater vs model", "ML rater vs model"],
      partB, [4.2, 3.4, 3.6, 3.4])

heading("C. ", "Aggregate figures per rater")
table(["Measure", "Clinical rater", "ML rater", "Model"], partC, [8.6, 2.6, 2.6, 2.6])

legend = doc.add_paragraph()
r = legend.add_run(
    "Second random 10% subsample of the included studies (manuscript references "
    + ", ".join(str(x) for x in refs) + "; Table S12 B). Two authors (clinical medicine, V.O.I.; machine "
    "learning, C.G.) rated independently from the same written guidance as the model, blind to its "
    "output, on the full QUADAS-2 (11 signalling questions, 3 applicability judgements) and 18 "
    "judgement-dependent CLAIM 2024 items (8, 11, 14 to 18, 20, 21, 26, 29, 30, 33, 36 to 39, 43). "
    "CLAIM responses yes / no / not applicable; QUADAS-2 signalling yes / no / unclear; applicability "
    "low / high / unclear. Compliance = yes / (yes + no). Domain-level risk derived from the signalling "
    "questions (any no: high; all yes: low; otherwise unclear). Gwet's AC1 is given alongside kappa "
    "because several items have a near-degenerate marginal distribution.")
r.italic = True
r.font.size = Pt(8)
r.font.name = "Arial"

doc.save(OUT_PATH)
print(f"Saved -> {OUT_PATH}")
