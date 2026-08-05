#!/usr/bin/env python3
"""
01_tab_1.py
Table 1 - Summary characteristics of included studies

Systematic Review: "Toward clinical-grade deep learning in coronary angiography"

Input:  xca_dl_review.db, xca_dl_review_references.bib
Output: results/tables/Table1_characteristics.docx

Publication period and publication type are derived from the bibliography;
study design is derived from the extraction database.
"""

import os
from collections import Counter

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

import common as C

# Publication periods come from the config file ([table1] periods).
PERIODS = C.PERIODS

OUT_PATH = os.path.join(C.OUT_TABLES, "Table1_characteristics.docx")

### Data
conn = C.connect()
papers = C.load_papers(conn)
N = len(papers)

year_counts = Counter(p.year for p in papers)
period_counts = [
    (lo, hi, sum(n for y, n in year_counts.items() if lo <= y <= hi))
    for lo, hi in PERIODS
]
assert sum(n for *_, n in period_counts) == N, "publication periods do not cover all studies"

n_journal = sum(1 for p in papers if p.is_journal)
n_conference = N - n_journal

# Study design: no included study reports prospective validation, so all are
# retrospective. Derived rather than asserted, so the claim tracks the data.
n_prospective = sum(
    1 for p in papers
    if C.scalar(conn, p.paper_id, "prospective_validation").lower() in ("y", "yes")
)
n_retrospective = N - n_prospective

min_year, max_year = min(year_counts), max(year_counts)
conn.close()


def period_label(lo, hi):
    """Collapse single-year bins (e.g. 2026-2026) to a bare year."""
    return f"{lo}" if lo == hi else f"{lo}–{hi}"

### Console summary
print(C.describe_config())
print(f"\nTable 1: Summary characteristics (N = {N})")
print(f"  Reference numbers: {papers[0].ref}\u2013{papers[-1].ref}")
print(f"  Publication period: {min_year}\u2013{max_year}")
for lo, hi, n in period_counts:
    print(f"    {period_label(lo, hi)}: {C.fmt(n, N)}")
print("  Publication type:")
print(f"    Journal article:  {C.fmt(n_journal, N)}")
print(f"    Conference paper: {C.fmt(n_conference, N)}")
print("  Study design:")
print(f"    Retrospective: {C.fmt(n_retrospective, N)}")
print(f"    Prospective:   {C.fmt(n_prospective, N)}")
print(f"  Per-year: {dict(sorted(year_counts.items()))}")

### DOCX
rows_data = [("Publication period", "", True)]
rows_data += [(f"    {period_label(lo, hi)}", C.fmt(n, N), False) for lo, hi, n in period_counts]
rows_data += [
    ("Publication type", "", True),
    ("    Journal article", C.fmt(n_journal, N), False),
    ("    Conference paper", C.fmt(n_conference, N), False),
    ("Study design", "", True),
    ("    Retrospective", C.fmt(n_retrospective, N), False),
    ("    Prospective", C.fmt(n_prospective, N), False),
]

C.ensure_dirs()
doc = Document()

section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = section.bottom_margin = Cm(2.5)
section.left_margin = section.right_margin = Cm(2.5)

p = doc.add_paragraph()
run = p.add_run("Table 1. ")
run.bold = True
run.font.size = Pt(11)
run.font.name = "Arial"
run = p.add_run(f"Summary characteristics of included studies (N = {N})")
run.font.size = Pt(11)
run.font.name = "Arial"

table = doc.add_table(rows=len(rows_data) + 1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
table.style = "Table Grid"

for i, text in enumerate(["Characteristic", "n (%)"]):
    cell = table.rows[0].cells[i]
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = True
    run.font.size = Pt(10)
    run.font.name = "Arial"
    tcPr = cell._element.get_or_add_tcPr()
    tcPr.append(tcPr.makeelement(
        qn("w:shd"), {qn("w:fill"): "E8E8E8", qn("w:val"): "clear"}))

for i, (label, value, is_header) in enumerate(rows_data):
    row = table.rows[i + 1]
    for j, (text, bold) in enumerate([(label, is_header), (value, False)]):
        cell = row.cells[j]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = bold
        run.font.size = Pt(10)
        run.font.name = "Arial"

for row in table.rows:
    row.cells[0].width = Cm(10)
    row.cells[1].width = Cm(4)

p = doc.add_paragraph()
run = p.add_run(
    "Publication periods correspond to three-year intervals chosen post hoc to "
    "summarise temporal trends in study output; they do not reflect predefined "
    "technological or methodological milestones. The final row (2026) is shown "
    "separately as the current, incomplete publication year. "
    "Publication type follows the venue recorded in the bibliography. "
    "Study design was classified as prospective only where validation was "
    "performed prospectively; no included study met this criterion."
)
run.italic = True
run.font.size = Pt(8)
run.font.name = "Arial"

doc.save(OUT_PATH)
print(f"\nSaved \u2192 {OUT_PATH}")
