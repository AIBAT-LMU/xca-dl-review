#!/usr/bin/env python3
"""
02_tab_s1.py
Table S1 - Application domains and reference lists

Systematic Review: "Toward clinical-grade deep learning in coronary angiography"

Input:  xca_dl_review.db
Output: results/tables/TableS1_applications.docx

Reports the same six domains as Figure 2B, each broken down into the
extraction sub-categories, with the reference numbers of the contributing
studies. Domain definitions come from common.DOMAIN_MAP, shared with
02_fig_2.py.
"""

import os

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

import common as C

OUT_PATH = os.path.join(C.OUT_TABLES, "TableS1_applications.docx")

### Data
conn = C.connect()
papers = C.load_papers(conn)
N = len(papers)
ref_of = {p.paper_id: p.ref for p in papers}

by_domain, by_sub = C.application_domains(conn, papers)
per_paper = C.domains_per_paper(by_domain, papers)
conn.close()

single = [pid for pid, d in per_paper.items() if len(d) == 1]
multi = [pid for pid, d in per_paper.items() if len(d) > 1]

# (label, n (%), references, indent, is_header)
rows = [
    ("Task structure", "", "", 0, True),
    ("Single domain", C.fmt(len(single), N),
     C.ref_ranges(ref_of[p] for p in single), 1, False),
    ("Multiple domains", C.fmt(len(multi), N),
     C.ref_ranges(ref_of[p] for p in multi), 1, False),
    ("Application domain", "", "", 0, True),
]

for domain in C.DOMAIN_ORDER:
    pids = by_domain.get(domain, set())
    subs = [s for s in C.SUBDOMAIN_ORDER[domain] if by_sub.get(s)]
    # A domain with a single sub-category carries its references directly.
    split = len(subs) > 1
    rows.append((domain, C.fmt(len(pids), N),
                 "" if split else C.ref_ranges(ref_of[p] for p in pids),
                 1, False))
    if split:
        for sub in subs:
            sp = by_sub[sub]
            rows.append((sub, C.fmt(len(sp), N),
                         C.ref_ranges(ref_of[p] for p in sp), 2, False))

### Console summary
print(f"Table S1: Application domains (N = {N})")
for label, n_str, _, indent, is_hdr in rows:
    if is_hdr:
        print(f"\n{label}")
    else:
        print(f"{'  ' * indent}{label}: {n_str}")

### DOCX
C.ensure_dirs()
doc = Document()

section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(2)
section.left_margin = section.right_margin = Cm(2)

p = doc.add_paragraph()
run = p.add_run("Table S1. ")
run.bold = True
run.font.size = Pt(11)
run.font.name = "Arial"
run = p.add_run("Application domains and reference lists")
run.font.size = Pt(11)
run.font.name = "Arial"

table = doc.add_table(rows=len(rows) + 1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
table.style = "Table Grid"

for i, text in enumerate(["Category", "n (%)", "References"]):
    cell = table.rows[0].cells[i]
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = True
    run.font.size = Pt(9)
    run.font.name = "Arial"
    tcPr = cell._element.get_or_add_tcPr()
    tcPr.append(tcPr.makeelement(
        qn("w:shd"), {qn("w:fill"): "E8E8E8", qn("w:val"): "clear"}))

for i, (label, n_str, refs, indent, is_hdr) in enumerate(rows):
    row = table.rows[i + 1]
    cell = row.cells[0]
    cell.text = ""
    run = cell.paragraphs[0].add_run("\u00A0\u00A0\u00A0\u00A0" * indent + label)
    run.bold = is_hdr
    run.italic = not is_hdr and indent == 2
    run.font.size = Pt(9)
    run.font.name = "Arial"
    for j, (text, size) in enumerate(((n_str, 9), (refs, 8))):
        cell = row.cells[j + 1]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.font.size = Pt(size)
        run.font.name = "Arial"

for row in table.rows:
    row.cells[0].width = Cm(5)
    row.cells[1].width = Cm(2)
    row.cells[2].width = Cm(10)

p = doc.add_paragraph()
run = p.add_run(
    "Reference numbers correspond to the supplementary reference list L1. "
    "Studies addressing multiple domains appear in multiple categories, so "
    "column totals exceed N. Percentages are of all included studies."
)
run.italic = True
run.font.size = Pt(8)
run.font.name = "Arial"

doc.save(OUT_PATH)
print(f"\nSaved \u2192 {OUT_PATH}")
